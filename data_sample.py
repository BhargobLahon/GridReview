import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_sample_workbook(filename="sample_results.xlsx"):
    # Industrial-grade contingency study dataset (IEEE Sub-transmission & Transmission benchmark style)
    data = [
        # N-0 Normal Operating Conditions
        ["N-0: Base Case Peak", "BUS-400-01 (Slack)", "Bus", 400.0, 1.025, None, "Yes"],
        ["N-0: Base Case Peak", "BUS-220-02 (PV Gen)", "Bus", 220.0, 1.040, None, "Yes"],
        ["N-0: Base Case Peak", "BUS-220-03 (Load)", "Bus", 220.0, 0.975, None, "Yes"],
        ["N-0: Base Case Peak", "LINE-400/220-Tr1", "Transformer", 400.0, None, 68.4, "Yes"],
        ["N-0: Base Case Peak", "LINE-220-01-04", "Line", 220.0, None, 74.2, "Yes"],
        ["N-0: Base Case Peak", "BUS-132-05 (Sub)", "Bus", 132.0, 0.960, None, "Yes"],

        # N-1 Contingency: Major Generation Trip
        ["N-1: Gen 2 Trip", "BUS-220-02 (PV Gen)", "Bus", 220.0, 0.910, None, "Yes"],
        ["N-1: Gen 2 Trip", "BUS-132-05 (Sub)", "Bus", 132.0, 0.885, None, "Yes"],
        ["N-1: Gen 2 Trip", "LINE-220-03-06", "Line", 220.0, None, 98.5, "Yes"],
        ["N-1: Gen 2 Trip", "LINE-132-05-08", "Line", 132.0, None, 104.2, "Yes"],

        # N-1 Contingency: Transmission Line Outage (Thermal Overload & Voltage Sag)
        ["N-1: Line 400-02 Outage", "BUS-220-04 (Load)", "Bus", 220.0, 0.920, None, "Yes"],
        ["N-1: Line 400-02 Outage", "LINE-400-03-05", "Line", 400.0, None, 114.5, "Yes"],
        ["N-1: Line 400-02 Outage", "TR-220/33-North", "Transformer", 220.0, None, 118.0, "Yes"],
        ["N-1: Line 400-02 Outage", "BUS-33-11 (Industrial)", "Bus", 33.0, 0.935, None, "Yes"],

        # Severe Contingency: Heavy Reactive Power Deficit (Critical Overload & Undervoltage)
        ["N-1: Heavy Corridor Outage", "BUS-132-09 (Load)", "Bus", 132.0, 0.860, None, "Yes"],
        ["N-1: Heavy Corridor Outage", "LINE-132-08-09", "Line", 132.0, None, 126.5, "Yes"],
        ["N-1: Heavy Corridor Outage", "TR-132/33-East", "Transformer", 132.0, None, 122.0, "Yes"],

        # Solver Non-Convergence Case (Voltage Collapse / Islanding Scenario)
        ["N-2: Dual Line Contingency", "SYSTEM-WIDE", "System", 400.0, None, None, "No"],

        # Light Load Condition (Overvoltage Risk)
        ["N-0: Base Case Off-Peak", "BUS-400-01 (Slack)", "Bus", 400.0, 1.055, None, "Yes"],
        ["N-0: Base Case Off-Peak", "BUS-220-03 (Load)", "Bus", 220.0, 1.065, None, "Yes"],
        ["N-0: Base Case Off-Peak", "LINE-220-02-04", "Line", 220.0, None, 45.0, "Yes"],
    ]

    columns = ["Case", "Element", "Type", "Nominal_kV", "Voltage_pu", "Loading_percent", "Converged"]
    df = pd.DataFrame(data, columns=columns)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Contingency_Results", startrow=3)
        ws = writer.sheets["Contingency_Results"]
        
        # Professional Industrial Header Metadata Block
        ws["A1"] = "PROJECT: Regional Transmission Network Interconnection Study (N-1/N-2 Security Assessment)"
        ws["A2"] = "CONFIDENTIAL & PROPRIETARY — SIMULATION DATASET EXCLUSIVE TO INTERNAL QA BENCHMARKING"
        ws["A1"].font = Font(bold=True, size=11, color="1F4E79")
        ws["A2"].font = Font(italic=True, size=9, color="7F7F7F")

    print(f"Generated industrial-grade dataset: {filename}")

if __name__ == "__main__":
    create_sample_workbook()