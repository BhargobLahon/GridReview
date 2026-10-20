import os
import streamlit as st
import pandas as pd
from openpyxl.styles import Font, PatternFill
from review_engine import (
    GRID_CODE_PRESETS,
    parse_nl_config,
    run_study_review,
    export_reviewed_excel,
    REQUIRED_COLUMNS,
)

def ensure_industrial_workbook(filename="sample_results.xlsx"):
    """Self-contained industrial-grade dataset generator (IEEE sub-transmission style)."""
    data = [
        ["N-0: Base Case Peak", "BUS-400-01 (Slack)", "Bus", 400.0, 1.025, None, "Yes"],
        ["N-0: Base Case Peak", "BUS-220-02 (PV Gen)", "Bus", 220.0, 1.040, None, "Yes"],
        ["N-0: Base Case Peak", "BUS-220-03 (Load)", "Bus", 220.0, 0.975, None, "Yes"],
        ["N-0: Base Case Peak", "LINE-400/220-Tr1", "Transformer", 400.0, None, 68.4, "Yes"],
        ["N-0: Base Case Peak", "LINE-220-01-04", "Line", 220.0, None, 74.2, "Yes"],
        ["N-0: Base Case Peak", "BUS-132-05 (Sub)", "Bus", 132.0, 0.960, None, "Yes"],
        ["N-1: Gen 2 Trip", "BUS-220-02 (PV Gen)", "Bus", 220.0, 0.910, None, "Yes"],
        ["N-1: Gen 2 Trip", "BUS-132-05 (Sub)", "Bus", 132.0, 0.885, None, "Yes"],
        ["N-1: Gen 2 Trip", "LINE-220-03-06", "Line", 220.0, None, 98.5, "Yes"],
        ["N-1: Gen 2 Trip", "LINE-132-05-08", "Line", 132.0, None, 104.2, "Yes"],
        ["N-1: Line 400-02 Outage", "BUS-220-04 (Load)", "Bus", 220.0, 0.920, None, "Yes"],
        ["N-1: Line 400-02 Outage", "LINE-400-03-05", "Line", 400.0, None, 114.5, "Yes"],
        ["N-1: Line 400-02 Outage", "TR-220/33-North", "Transformer", 220.0, None, 118.0, "Yes"],
        ["N-1: Line 400-02 Outage", "BUS-33-11 (Industrial)", "Bus", 33.0, 0.935, None, "Yes"],
        ["N-1: Heavy Corridor Outage", "BUS-132-09 (Load)", "Bus", 132.0, 0.860, None, "Yes"],
        ["N-1: Heavy Corridor Outage", "LINE-132-08-09", "Line", 132.0, None, 126.5, "Yes"],
        ["N-1: Heavy Corridor Outage", "TR-132/33-East", "Transformer", 132.0, None, 122.0, "Yes"],
        ["N-2: Dual Line Contingency", "SYSTEM-WIDE", "System", 400.0, None, None, "No"],
        ["N-0: Base Case Off-Peak", "BUS-400-01 (Slack)", "Bus", 400.0, 1.055, None, "Yes"],
        ["N-0: Base Case Off-Peak", "BUS-220-03 (Load)", "Bus", 220.0, 1.065, None, "Yes"],
        ["N-0: Base Case Off-Peak", "LINE-220-02-04", "Line", 220.0, None, 45.0, "Yes"],
    ]
    columns = ["Case", "Element", "Type", "Nominal_kV", "Voltage_pu", "Loading_percent", "Converged"]
    df = pd.DataFrame(data, columns=columns)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Contingency_Results", startrow=3)
        ws = writer.sheets["Contingency_Results"]
        ws["A1"] = "PROJECT: Regional Transmission Network Interconnection Study (N-1/N-2 Security Assessment)"
        ws["A2"] = "CONFIDENTIAL & PROPRIETARY — SIMULATION DATASET EXCLUSIVE TO INTERNAL QA BENCHMARKING"
        ws["A1"].font = Font(bold=True, size=11, color="1F4E79")
        ws["A2"].font = Font(italic=True, size=9, color="7F7F7F")

st.set_page_config(page_title="StudySense: Automated Power-System Study Results Reviewer", layout="wide")

st.title("StudySense: Automated Power-System Study Results Reviewer")
st.caption("Engineering Automation: Natural-Language Configuration + Dynamic Column Mapping + N-0/N-1 Validation")

# Sidebar Configuration
with st.sidebar:
    st.header("1. Grid Code Preset")
    profile_options = list(GRID_CODE_PRESETS.keys()) + ["+ Create Custom Profile..."]
    selected_profile = st.selectbox("Select Criteria Profile", profile_options)

    if selected_profile == "+ Create Custom Profile...":
        custom_name = st.text_input("Profile Name", value="Custom Utility Standard")
        c_vmin = st.number_input("Custom Min Voltage (pu)", 0.80, 1.00, 0.95, 0.01)
        c_vmax = st.number_input("Custom Max Voltage (pu)", 1.00, 1.20, 1.05, 0.01)
        c_load = st.number_input("Custom Loading Limit (%)", 80.0, 130.0, 100.0, 5.0)
        
        if st.button("Save Custom Profile"):
            GRID_CODE_PRESETS[custom_name] = {
                "v_min_pu": c_vmin,
                "v_max_pu": c_vmax,
                "v_crit_min_pu": c_vmin - 0.05,
                "v_crit_max_pu": c_vmax + 0.05,
                "v_warn_margin_pu": 0.01,
                "load_warn_pct": c_load - 10.0,
                "load_limit_n0_pct": c_load,
                "load_limit_n1_pct": c_load + 10.0,
                "load_crit_pct": c_load + 20.0,
            }
            st.success(f"Profile '{custom_name}' activated!")
            st.rerun()
    else:
        if "active_cfg" not in st.session_state or st.session_state.get("last_preset") != selected_profile:
            st.session_state.active_cfg = GRID_CODE_PRESETS[selected_profile].copy()
            st.session_state.last_preset = selected_profile

    differentiate_n1 = st.checkbox("Apply separate N-1 Emergency Loading Limit", value=True)

    st.header("2. Editable Thresholds")
    cfg = st.session_state.active_cfg
    cfg["v_min_pu"] = st.number_input("Min Voltage Limit (pu)", 0.80, 1.00, float(cfg["v_min_pu"]), 0.01)
    cfg["v_max_pu"] = st.number_input("Max Voltage Limit (pu)", 1.00, 1.20, float(cfg["v_max_pu"]), 0.01)
    cfg["v_crit_min_pu"] = st.number_input("Critical Undervoltage (pu)", 0.70, 0.95, float(cfg["v_crit_min_pu"]), 0.01)
    cfg["load_warn_pct"] = st.number_input("Loading Warning (%)", 50.0, 100.0, float(cfg["load_warn_pct"]), 5.0)
    cfg["load_limit_n0_pct"] = st.number_input("N-0 Loading Violation (%)", 80.0, 120.0, float(cfg["load_limit_n0_pct"]), 5.0)
    if differentiate_n1:
        cfg["load_limit_n1_pct"] = st.number_input("N-1 Emergency Loading Limit (%)", 90.0, 130.0, float(cfg["load_limit_n1_pct"]), 5.0)
    cfg["load_crit_pct"] = st.number_input("Critical Loading (%)", 100.0, 150.0, float(cfg["load_crit_pct"]), 5.0)

# Natural-Language Threshold Configuration Bar
with st.expander("Natural-Language Threshold Configuration (Human-in-the-Loop)", expanded=True):
    nl_prompt = st.text_input(
        "Instruct the configuration agent in plain English:",
        value="Review this study using voltage limits of 0.95 to 1.05 pu. Treat loading above 100% as a violation and loading above 90% as a warning."
    )
    if st.button("Parse Instruction & Update Thresholds"):
        new_cfg, trace_notes = parse_nl_config(nl_prompt, st.session_state.active_cfg)
        st.session_state.active_cfg = new_cfg
        st.session_state.nl_trace = trace_notes
        st.rerun()

    if "nl_trace" in st.session_state:
        st.json(st.session_state.active_cfg)

# File Upload or Built-in Industrial Sample
col_up, col_demo = st.columns([2, 1])
with col_up:
    uploaded_file = st.file_uploader("Upload Custom Study Results (.xlsx or .csv)", type=["xlsx", "csv"])
with col_demo:
    st.write("Quick Benchmark Demo:")
    use_sample = st.button("Load Industrial Substation Dataset")

df_input = None
df_raw = None

if uploaded_file:
    if uploaded_file.name.endswith(".csv"):
        df_raw = pd.read_csv(uploaded_file)
    else:
        df_temp = pd.read_excel(uploaded_file)
        if not REQUIRED_COLUMNS.issubset(set(df_temp.columns)) and len(df_temp) > 3:
            df_raw = pd.read_excel(uploaded_file, skiprows=3)
        else:
            df_raw = df_temp
elif use_sample or st.session_state.get("sample_loaded"):
    st.session_state.sample_loaded = True
    if not os.path.exists("sample_results.xlsx"):
        ensure_industrial_workbook("sample_results.xlsx")
    df_raw = pd.read_excel("sample_results.xlsx", skiprows=3)

if df_raw is not None:
    # Dynamic Column Mapping if columns don't match standard schema
    if REQUIRED_COLUMNS.issubset(set(df_raw.columns)):
        df_input = df_raw
    else:
        st.warning("⚠️ Uploaded file columns do not match standard schema. Please map your columns below:")
        cols = list(df_raw.columns)
        
        m_col1, m_col2 = st.columns(2)
        with m_col1:
            map_case = st.selectbox("Select Case / Contingency Column", cols, index=0 if "Case" in cols else 0)
            map_elem = st.selectbox("Select Element / Bus / Line Column", cols, index=0 if "Element" in cols else 0)
            map_type = st.selectbox("Select Equipment Type Column (Bus/Line/Transformer)", cols, index=0 if "Type" in cols else 0)
        with m_col2:
            map_volt = st.selectbox("Select Voltage (pu) Column", cols, index=0 if "Voltage_pu" in cols else 0)
            map_load = st.selectbox("Select Loading (%) Column", cols, index=0 if "Loading_percent" in cols else 0)
            map_conv = st.selectbox("Select Convergence Status Column (Yes/No)", cols, index=0 if "Converged" in cols else 0)

        if st.button("Apply Column Mapping & Run Review"):
            rename_dict = {
                map_case: "Case",
                map_elem: "Element",
                map_type: "Type",
                map_volt: "Voltage_pu",
                map_load: "Loading_percent",
                map_conv: "Converged"
            }
            df_input = df_raw.rename(columns=rename_dict)
            if "Nominal_kV" not in df_input.columns:
                df_input["Nominal_kV"] = 220.0
            st.success("Columns mapped successfully! Review initialized.")

if df_input is not None:
    try:
        results = run_study_review(df_input, st.session_state.active_cfg, differentiate_n1)
        m = results["summary_metrics"]
        rev_df = results["reviewed_df"]

        st.divider()
        c1, c2, c3, c4, c5 = st.columns(5)
        c1.metric("Cases Reviewed", m["total_cases"])
        c2.metric("Converged Cases", m["converged_cases"])
        c3.metric("Non-Convergent", m["non_converged_cases"])
        c4.metric("Voltage Violations", m["voltage_violations"])
        c5.metric("Equipment Overloads", m["equipment_overloads"])

        with st.expander("Execution Trace & Verification Log", expanded=False):
            st.markdown(f"""
            1. **[Schema Guard]** Verified required columns or dynamic mapping across **{len(df_input)}** element rows.
            2. **[Contingency Classifier]** Tagged N-0 Normal vs. N-1/N-2 Contingency regimes (`differentiate_n1={differentiate_n1}`).
            3. **[Convergence Short-Circuit]** Isolated **{m['non_converged_cases']}** non-convergent case(s) prior to floating-point limit checks.
            4. **[Severity Auditor]** Ranked findings into `CRITICAL`, `VIOLATION`, `WARNING`, and `PASS`.
            """)

        tab1, tab2, tab3 = st.tabs(["Ranked Findings (All)", "Critical & Violations Only", "Automated Engineering Report"])
        with tab1:
            st.dataframe(rev_df, use_container_width=True)
        with tab2:
            st.dataframe(rev_df[rev_df["Severity"].isin(["CRITICAL", "VIOLATION", "WARNING"])], use_container_width=True)
        with tab3:
            st.code(results["text_report"], language="text")

        dl1, dl2 = st.columns(2)
        with dl1:
            excel_bytes = export_reviewed_excel(rev_df)
            st.download_button(
                "Download Color-Coded Reviewed Excel (.xlsx)",
                data=excel_bytes,
                file_name="StudySense_Audited_Results.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        with dl2:
            st.download_button(
                "Download Executive Summary (.txt)",
                data=results["text_report"],
                file_name="StudySense_Executive_Summary.txt",
                mime="text/plain"
            )
    except Exception as e:
        st.error(f"Validation Error: {e}")